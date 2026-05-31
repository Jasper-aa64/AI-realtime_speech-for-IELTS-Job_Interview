#!/usr/bin/env python3
from __future__ import annotations

import re
import shutil
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK, WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
DOC_DIR = ROOT / "软件构造"
TEMPLATE = DOC_DIR / "软件构造_实验1_报告.docx"
FIG_DIR = DOC_DIR / "generated_figures"
FIG_DIR.mkdir(exist_ok=True)

FONT_HEI = "/System/Library/Fonts/STHeiti Medium.ttc"
FONT_SONG = "/System/Library/Fonts/Supplemental/Songti.ttc"


def font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    path = FONT_HEI if bold else FONT_SONG
    return ImageFont.truetype(path, size)


def draw_wrapped_text(draw: ImageDraw.ImageDraw, xy: tuple[int, int], text: str, fnt, fill=(20, 30, 45), width=260, line_gap=6):
    x, y = xy
    line = ""
    lines: list[str] = []
    for ch in text:
        trial = line + ch
        if draw.textbbox((0, 0), trial, font=fnt)[2] <= width:
            line = trial
        else:
            if line:
                lines.append(line)
            line = ch
    if line:
        lines.append(line)
    for line in lines:
        draw.text((x, y), line, font=fnt, fill=fill)
        y += fnt.size + line_gap
    return y


def box(draw, xy, text, *, fill="#F8FBFF", outline="#365F91", title=False, width=260):
    x1, y1, x2, y2 = xy
    draw.rounded_rectangle(xy, radius=12, fill=fill, outline=outline, width=3)
    fnt = font(23 if title else 20, bold=title)
    draw_wrapped_text(draw, (x1 + 18, y1 + 18), text, fnt, width=width)


def arrow(draw, start, end, color="#4A5568"):
    draw.line([start, end], fill=color, width=4)
    ex, ey = end
    sx, sy = start
    if ex >= sx:
        pts = [(ex, ey), (ex - 14, ey - 8), (ex - 14, ey + 8)]
    else:
        pts = [(ex, ey), (ex + 14, ey - 8), (ex + 14, ey + 8)]
    draw.polygon(pts, fill=color)


def save_diagram(name: str, title: str, boxes: list[tuple[int, int, int, int, str]], arrows: list[tuple[tuple[int, int], tuple[int, int]]] = ()):
    img = Image.new("RGB", (1400, 820), "#FFFFFF")
    draw = ImageDraw.Draw(img)
    draw.text((40, 32), title, font=font(34, bold=True), fill="#111827")
    draw.line([(40, 86), (1360, 86)], fill="#365F91", width=4)
    for item in boxes:
        x1, y1, x2, y2, text = item
        box(draw, (x1, y1, x2, y2), text, width=max(180, x2 - x1 - 36))
    for s, e in arrows:
        arrow(draw, s, e)
    path = FIG_DIR / f"{name}.png"
    img.save(path)
    return path


FIGURES = {
    "exp2_patterns": save_diagram(
        "exp2_patterns",
        "实验2：AI Provider 四种设计模式协同",
        [
            (70, 150, 350, 300, "Strategy\nAIProvider 统一不同 AI 后端"),
            (390, 150, 670, 300, "Adapter\nCodexCliClient / HttpApiProvider 适配外部调用"),
            (710, 150, 990, 300, "Chain of Responsibility\nHTTP → Codex → Fallback"),
            (1030, 150, 1310, 300, "Template Method\nAiTaskTemplate 固定任务骨架"),
            (390, 470, 990, 650, "统一结果对象 ProviderRunResult\n成功、失败、跳过、fallback 状态透明记录"),
        ],
        [((350, 225), (390, 225)), ((670, 225), (710, 225)), ((990, 225), (1030, 225)), ((670, 300), (670, 470))],
    ),
    "exp2_chain": save_diagram(
        "exp2_chain",
        "实验2：Provider 责任链与透明降级",
        [
            (80, 190, 330, 330, "业务任务\n写作评分 / 口语报告 / 追问"),
            (410, 190, 660, 330, "HTTP Provider\nAiapis / OpenAI 兼容接口"),
            (740, 190, 990, 330, "Codex Provider\nCLI 兼容旧链路"),
            (1070, 190, 1320, 330, "Fallback Provider\n只能标记 fallback，不伪装成功"),
            (410, 520, 990, 670, "ProviderRunResult\n保留 backend / provider / usage / error 证据"),
        ],
        [((330, 260), (410, 260)), ((660, 260), (740, 260)), ((990, 260), (1070, 260)), ((700, 330), (700, 520))],
    ),
    "exp3_tdd": save_diagram(
        "exp3_tdd",
        "实验3：Red - Green - Refactor 循环",
        [
            (120, 180, 420, 330, "Red\n先写失败测试，暴露接口契约"),
            (550, 180, 850, 330, "Green\n最小实现让测试通过"),
            (980, 180, 1280, 330, "Refactor\n在测试保护下整理结构"),
            (390, 530, 1010, 680, "验证脚本\nC++ audio_core / WASM / Django AI provider / metrics"),
        ],
        [((420, 255), (550, 255)), ((850, 255), (980, 255)), ((1130, 330), (1010, 530)), ((390, 605), (270, 330))],
    ),
    "exp3_layers": save_diagram(
        "exp3_layers",
        "实验3：测试分层",
        [
            (110, 160, 420, 300, "C++ 单元测试\naudio_core_test 固定 RMS、VAD、裁剪"),
            (545, 160, 855, 300, "WASM 边界测试\nmetrics 与 frame summary"),
            (980, 160, 1290, 300, "Django 回归测试\napps.ai / speaking provider"),
            (350, 500, 1050, 650, "一键验证脚本\n把跨语言测试收束为可重复执行的质量门"),
        ],
        [((265, 300), (520, 500)), ((700, 300), (700, 500)), ((1135, 300), (880, 500))],
    ),
    "exp4_before_after": save_diagram(
        "exp4_before_after",
        "实验4：重构前后结构对比",
        [
            (80, 150, 620, 330, "重构前\n采集、设备 IO、协议、VAD、业务状态混在一起\n测试依赖设备和网络"),
            (780, 150, 1320, 330, "重构后\n抽出 audio_core 纯核\nIO、WASM、业务编排只依赖统一分析结果"),
            (170, 530, 1230, 670, "核心变化：Extract Module + Strategy + Adapter + Characterization Tests"),
        ],
        [((620, 240), (780, 240)), ((700, 330), (700, 530))],
    ),
    "exp4_layers": save_diagram(
        "exp4_layers",
        "实验4：audio_core 分层依赖",
        [
            (110, 130, 390, 270, "浏览器录音\nMediaRecorder / AudioWorklet"),
            (530, 130, 870, 270, "WASM Adapter\naudio_core_wasm.cpp"),
            (1010, 130, 1290, 270, "Django 业务\nP1/P2/P3 baseline"),
            (380, 470, 1020, 650, "C++ audio_core 纯核\nRMS / Peak / Trim / Resample / libfvad VAD"),
        ],
        [((390, 200), (530, 200)), ((870, 200), (1010, 200)), ((700, 270), (700, 470))],
    ),
}


def remove_after_cover(doc: Document):
    body = doc.element.body
    keep = []
    table_seen = False
    for child in list(body):
        keep.append(child)
        if child.tag.endswith("tbl"):
            table_seen = True
            break
    if not table_seen:
        raise RuntimeError("template cover table not found")
    for child in list(body):
        if child not in keep and not child.tag.endswith("sectPr"):
            body.remove(child)


def set_cell_text(cell, text: str, size: int = 14, bold: bool = False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.bold = bold
    r.font.name = "Times New Roman"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    r.font.size = Pt(size)


def update_cover(doc: Document, number: str, name: str, date: str = "2026.5.31"):
    table = doc.tables[0]
    values = {
        0: ("实验课程：", "软件构造"),
        1: ("实验编号：", number),
        2: ("实验名称：", name),
        3: ("实验人员：", "学号", "23111302079"),
        4: ("实验人员：", "姓名", "梁峻铭"),
        5: ("实验人员：", "班级", "23软件工程"),
        6: ("指导教师：", "李文杰"),
        7: ("实验室：", "2060201"),
        8: ("实验日期：", date),
    }
    for idx, row in enumerate(table.rows):
        vals = values[idx]
        if len(vals) == 2:
            set_cell_text(row.cells[0], vals[0], 14, True)
            set_cell_text(row.cells[1], vals[1], 13)
            set_cell_text(row.cells[2], vals[1], 13)
        else:
            set_cell_text(row.cells[0], vals[0], 14, True)
            set_cell_text(row.cells[1], vals[1], 13)
            set_cell_text(row.cells[2], vals[2], 13)


def set_run_font(run, size=12, bold=False, color=None):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_heading(doc: Document, title: str):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(title)
    set_run_font(r, 15, True)
    p_pr = p._p.get_or_add_pPr()
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "3")
    bottom.set(qn("w:color"), "365F91")
    border.append(bottom)
    p_pr.append(border)


def add_subheading(doc: Document, title: str):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(title)
    set_run_font(r, 13, True)


def add_para(doc: Document, text: str):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Pt(24)
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    set_run_font(r, 11)


def add_bullet(doc: Document, text: str):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Pt(18)
    p.paragraph_format.first_line_indent = Pt(-10)
    p.paragraph_format.line_spacing = 1.2
    r = p.add_run("• " + text)
    set_run_font(r, 11)


def add_caption(doc: Document, text: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text)
    set_run_font(r, 10, False, (80, 80, 80))


def add_figure(doc: Document, path: Path, caption: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Inches(5.9))
    add_caption(doc, caption)


def add_table(doc: Document, headers: list[str], rows: list[list[str]]):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_text(cell, h, 10, True)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value, 10, False)
    doc.add_paragraph()


def plain_lines(md: str) -> list[str]:
    out = []
    for line in md.splitlines():
        line = line.strip()
        if not line or line.startswith(">"):
            continue
        line = re.sub(r"`([^`]+)`", r"\1", line)
        line = line.replace("**", "").replace("|", " / ")
        if line.startswith("#"):
            continue
        out.append(line)
    return out


def add_selected_source(doc: Document, md_path: Path, limit: int = 14):
    lines = plain_lines(md_path.read_text(encoding="utf-8"))
    picked = [line for line in lines if len(line) > 18][:limit]
    for line in picked:
        if line.startswith("- "):
            add_bullet(doc, line[2:])
        else:
            add_para(doc, line)


def build_report(number: str, name: str, sections: list[tuple[str, list[tuple[str, object]]]], figures: list[tuple[Path, str]], out_path: Path):
    doc = Document(str(TEMPLATE))
    remove_after_cover(doc)
    update_cover(doc, number, name)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    figure_index = 0
    for title, blocks in sections:
        add_heading(doc, title)
        for kind, value in blocks:
            if kind == "p":
                add_para(doc, str(value))
            elif kind == "b":
                add_bullet(doc, str(value))
            elif kind == "sub":
                add_subheading(doc, str(value))
            elif kind == "table":
                headers, rows = value
                add_table(doc, headers, rows)
            elif kind == "source":
                path, limit = value
                add_selected_source(doc, path, limit)
            elif kind == "fig":
                fig_path, caption = figures[figure_index]
                figure_index += 1
                add_figure(doc, fig_path, caption)
    doc.save(out_path)


def repack_existing_report_cover(number: str, name: str, path: Path):
    """Keep an existing report body, but replace its cover with the template cover."""
    if not path.exists():
        return

    source = Document(str(path))
    cover_seen = False
    body_content = []
    for child in list(source._element.body):
        tag = child.tag.split("}")[-1]
        if not cover_seen:
            if tag == "tbl":
                cover_seen = True
            continue
        if tag == "sectPr":
            continue
        if not body_content and tag == "p":
            text = "".join(t.text or "" for t in child.iter(qn("w:t"))).strip()
            if not text:
                continue
        body_content.append(deepcopy(child))

    doc = Document(str(TEMPLATE))
    remove_after_cover(doc)
    update_cover(doc, number, name)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    body = doc._element.body
    sect_pr = body.sectPr
    insert_at = list(body).index(sect_pr) if sect_pr is not None else len(body)
    for offset, element in enumerate(body_content):
        body.insert(insert_at + offset, element)
    doc.save(path)


def exp2_sections():
    src = DOC_DIR / "实验2_设计文档.md"
    return [
        ("一、实验目的", [
            ("p", "本实验以 IELTS Web 产品的 AI Provider 链路为对象，使用设计模式对已有代码进行结构化整理。目标不是改变产品功能，而是在已有 Django 后端基础上，把评分、追问、报告等 AI 调用抽象成可切换、可测试、可降级的架构。"),
            ("b", "掌握 Strategy、Adapter、Chain of Responsibility、Template Method 四种设计模式在真实项目中的落地方式。"),
            ("b", "理解外部模型 API、Codex CLI、fallback 兜底之间的职责边界。"),
            ("b", "形成类图、接口和关键代码证据，为后续 TDD、重构、开源复用实验奠定结构基础。"),
        ]),
        ("二、实验环境", [
            ("table", (["项目", "内容"], [
                ["开发语言", "Python 3 / Django；C++17 接口设计作为实时语音扩展方向"],
                ["主要模块", "backend_django/apps/ai、backend_django/apps/speaking、backend_django/apps/writing"],
                ["验证方式", "Django 单元测试、Provider 路由诊断脚本、类图与文档审查"],
                ["运行环境", "macOS 本地开发环境，SQLite 测试库，OpenAI-compatible HTTP provider 可配置"],
            ])),
        ]),
        ("三、实验内容与业务背景", [
            ("source", (src, 7)),
            ("fig", None),
        ]),
        ("四、设计模式实现说明", [
            ("sub", "1. Strategy：Provider 可运行时切换"),
            ("p", "AIProvider 抽象把不同模型后端视为可替换策略。HTTP provider、Codex provider、Mock/Fallback provider 都通过统一接口返回 ProviderRunResult，使业务代码不需要关心具体模型调用方式。"),
            ("sub", "2. Adapter：统一外部后端与外部进程"),
            ("p", "HttpApiProvider 适配 OpenAI-compatible Chat Completions；CodexCliClient 适配本地 codex exec。两者输入输出格式不同，但经 Adapter 层转换后，向业务层暴露一致结果。"),
            ("sub", "3. Chain of Responsibility：透明降级链"),
            ("p", "ProviderChain 让 HTTP、Codex、Fallback 按顺序处理任务。关键点是 fallback 不允许伪装成 SUCCESS，而是必须携带明确的 backend/status/error 元数据。"),
            ("sub", "4. Template Method：AI 任务执行骨架"),
            ("p", "AiTaskTemplate 固定 claim、build prompt、call provider、apply result、record status 的流程，各子类只实现差异步骤，减少写作评分和口语报告中的重复。"),
            ("fig", None),
        ]),
        ("五、类结构与协作关系", [
            ("source", (src, 10)),
            ("table", (["类/接口", "模式", "职责"], [
                ["AIProvider", "Strategy", "定义模型调用策略统一接口"],
                ["HttpApiProvider / CodexCliClient", "Adapter", "屏蔽 HTTP API 与 CLI 调用差异"],
                ["ProviderChain", "Chain of Responsibility", "组织 HTTP → Codex → Fallback 的降级链"],
                ["AiTaskTemplate", "Template Method", "复用 AI task 生命周期骨架"],
            ])),
        ]),
        ("六、测试与验证", [
            ("p", "实验中保留既有测试作为重构保护网，并补充 HTTP provider 与 follow-up routing 相关测试。验收重点不是只看类名是否出现，而是确认 provider 可切换、fallback 不伪装、usage 元数据可用于真实路由验收。"),
            ("b", "相关测试覆盖 provider 结果应用、HTTP provider 解析、P1/P3 follow-up 路由、streaming follow-up 路径。"),
            ("b", "真实 Aiapis 路由通过脚本验证 backend/provider/usage token，而不是只看 model 字段。"),
        ]),
        ("七、实验结论与心得", [
            ("p", "本实验把原来散落在业务服务中的 AI 调用逻辑抽象成可替换策略、适配层、责任链和模板流程。设计模式不是额外装饰，而是解决了外部模型不稳定、调用方式多样、fallback 需要透明记录等真实问题。"),
        ]),
    ], [(FIGURES["exp2_patterns"], "图 1 AI Provider 四种设计模式协同图"), (FIGURES["exp2_chain"], "图 2 Provider 责任链与透明降级图")]


def exp3_sections():
    src = DOC_DIR / "实验3_TDD验证报告.md"
    return [
        ("一、实验目的", [
            ("p", "本实验使用测试驱动开发思想，为 IELTS Web 产品的重构线建立可重复执行的质量保护网。目标是先用测试固定关键契约，再允许后续继续重构 audio_core、WASM、AI Provider 与指标采集层。"),
            ("b", "理解 Red-Green-Refactor 循环在真实项目中的作用。"),
            ("b", "为跨 C++、WASM、Django 的重构建立分层测试。"),
            ("b", "把人工判断转化为可执行的一键验证脚本。"),
        ]),
        ("二、实验环境", [
            ("table", (["项目", "内容"], [
                ["C++ 测试", "tests/audio_core_test.cpp，验证 RMS、Peak、VAD、裁剪、重采样"],
                ["浏览器/WASM 测试", "scripts/test_wasm_audio_preprocessor_metrics.mjs"],
                ["Django 测试", "apps.ai、apps.speaking 等回归测试"],
                ["验证脚本", "scripts/run_experiment3_tdd_validation.sh"],
            ])),
        ]),
        ("三、TDD 测试对象", [
            ("source", (src, 9)),
            ("fig", None),
        ]),
        ("四、Red / Green / Refactor 过程", [
            ("sub", "1. Red：先暴露风险"),
            ("p", "在重构 audio_core 与 provider 链路前，先列出最容易破坏的契约：非法 frame 不能被静默接受、fallback 不能伪装成功、metrics 不能把大段音频塞入业务 payload。"),
            ("sub", "2. Green：最小实现通过测试"),
            ("p", "补齐测试和脚本后，先让核心路径通过，再继续扩展。这样每一步都能知道是新改动引入了问题，还是旧问题本来存在。"),
            ("sub", "3. Refactor：在测试保护下整理结构"),
            ("p", "测试通过后，再把 audio_core、WASM wrapper、AI provider 结构继续拆分，避免为了重构而破坏已有 P1/P2/P3 功能。"),
            ("fig", None),
        ]),
        ("五、测试用例设计", [
            ("table", (["测试层", "代表用例", "防止的回退"], [
                ["C++ audio_core", "固定样本验证 RMS、Peak、TrimSilence", "算法重构后数值漂移"],
                ["libfvad backend", "合法静音帧、非法 sample rate、非法 frame 长度", "开源库 wrapper 漏校验"],
                ["WASM metrics", "frame summary、preprocessor metrics", "前端 payload 膨胀"],
                ["Django provider", "HTTP / Codex / fallback 路由测试", "fallback 冒充 AI 成功"],
            ])),
            ("source", (src, 7)),
        ]),
        ("六、验证结果", [
            ("p", "实验报告中的验证结果以现有脚本和测试输出为准。若某些真实端点需要 key 或真人设备，则在报告中保留待测边界，不把自动化结果冒充成真人验收。"),
            ("source", (src, 5)),
        ]),
        ("七、实验结论与心得", [
            ("p", "TDD 在本项目中的价值不是增加测试数量，而是把重构中最容易出错的跨层契约固定下来。它让后续实验4重构和实验5复用不再依赖记忆和手工点击，而是有可以反复运行的质量门。"),
        ]),
    ], [(FIGURES["exp3_tdd"], "图 1 Red-Green-Refactor 循环图"), (FIGURES["exp3_layers"], "图 2 跨层测试分层图")]


def exp4_sections():
    src = DOC_DIR / "实验4_重构报告.md"
    return [
        ("一、实验目的", [
            ("p", "本实验针对 IELTS Web 产品中音频预处理和实时语音扩展相关的坏味道进行重构。目标是在不改变 P1/P2/P3、Mock、写作和报告等现有用户功能的前提下，把可复用的音频处理逻辑抽成纯核。"),
            ("b", "识别 IO、协议、算法、业务状态混杂带来的坏味道。"),
            ("b", "使用 Extract Module、Strategy、Adapter、Characterization Tests 等重构手法。"),
            ("b", "让同一套 C++ audio_core 能被 native 测试、WASM demo 和浏览器预处理复用。"),
        ]),
        ("二、实验环境", [
            ("table", (["项目", "内容"], [
                ["语言与工具", "C++17、CMake、Emscripten、Django、JavaScript"],
                ["重构对象", "src/ielts/audio_core.cpp、include/ielts/audio_core.h、wasm wrapper、前端预处理 runtime"],
                ["验证方式", "C++ tests、WASM 构建、Django check、脚本化验证"],
                ["边界", "不重写成熟产品功能，只重构内部结构和可选加速层"],
            ])),
        ]),
        ("三、重构前的问题", [
            ("source", (src, 9)),
            ("fig", None),
        ]),
        ("四、重构后的结构", [
            ("sub", "1. 抽出 audio_core 纯核"),
            ("p", "重构后，RMS、Peak、TrimSilence、Resample、VAD 分析都集中在 audio_core 中。该模块不依赖 Qt、PortAudio、Django 或浏览器对象，因此可以独立编译和测试。"),
            ("sub", "2. 用 Strategy 隔离 VAD backend"),
            ("p", "VadBackend::RmsThreshold 作为确定性 fallback，VadBackend::WebRtc 通过 libfvad 复用 WebRTC VAD。业务层只关心统一的 FrameAnalysis 结果。"),
            ("sub", "3. 用 Adapter 暴露 WASM 边界"),
            ("p", "audio_core_wasm.cpp 把 C++ 纯核转换为浏览器可调用的 C ABI，使前端 AudioWorklet 或预处理 runtime 可以在不理解 C++ 内部结构的情况下获得分析结果。"),
            ("fig", None),
        ]),
        ("五、使用的重构手法", [
            ("table", (["重构手法", "落点", "效果"], [
                ["Extract Module", "抽出 audio_core", "算法核可独立编译、测试、复用"],
                ["Extract Interface", "AudioFrameAnalyzer / VAD backend", "隔离业务与算法实现"],
                ["Strategy", "RmsThreshold / WebRtc backend", "方便替换算法与降级"],
                ["Adapter", "WASM wrapper", "把 C++ 纯核接入浏览器"],
                ["Characterization Tests", "audio_core_test 与验证脚本", "锁住重构前后行为"],
            ])),
        ]),
        ("六、关键代码证据与验证", [
            ("source", (src, 9)),
            ("p", "验证重点是确认重构后 baseline 行为保持可用，WASM 作为可选加速层出现，失败时仍能回到浏览器原生录音/批处理链路。"),
        ]),
        ("七、重构收益", [
            ("p", "重构后，可测试性、可复用性、可部署性和风险控制都得到提升。audio_core 可以被 native test、WASM demo、未来 realtime gateway 共用；前端正式业务流程仍保留 baseline，降低引入 C++/WASM 的风险。"),
            ("source", (src, 6)),
        ]),
        ("八、实验结论与心得", [
            ("p", "本次重构说明，生产系统中的重构不应以推倒重来为目标，而应围绕可验证边界逐步抽核、隔离和替换。audio_core 的拆分让实验5开源复用和后续实时语音扩展都有了稳定地基。"),
        ]),
    ], [(FIGURES["exp4_before_after"], "图 1 重构前后结构对比图"), (FIGURES["exp4_layers"], "图 2 audio_core 分层依赖图")]


def main():
    reports = [
        ("2", "利用设计模式进行类设计并编程实现", *exp2_sections(), DOC_DIR / "软件构造_实验2_报告.docx"),
        ("3", "测试驱动开发(TDD)", *exp3_sections(), DOC_DIR / "软件构造_实验3_报告.docx"),
        ("4", "软件重构", *exp4_sections(), DOC_DIR / "软件构造_实验4_报告.docx"),
    ]
    for number, name, sections, figures, out_path in reports:
        build_report(number, name, sections, figures, out_path)
        print(out_path)
    exp5_path = DOC_DIR / "软件构造_实验5_报告.docx"
    repack_existing_report_cover("5", "复用的软件资源进行软件构造", exp5_path)
    if exp5_path.exists():
        print(exp5_path)


if __name__ == "__main__":
    main()
